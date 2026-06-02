import asyncio
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional
from loguru import logger

from app.config import get_settings
from app.models import Document, DocumentSection, KnowledgeUnit, ExtractedEntity, DocumentStatus

settings = get_settings()


async def run_ingestion_pipeline(doc_id: str, file_path: str):
    from app.database import AsyncSessionLocal
    from sqlalchemy import select

    async with AsyncSessionLocal() as db:
        try:
            result = await db.execute(select(Document).where(Document.id == doc_id))
            doc = result.scalar_one_or_none()
            if not doc:
                logger.error(f"Document {doc_id} not found")
                return

            doc.status = DocumentStatus.PROCESSING
            doc.processing_started_at = datetime.utcnow()
            await db.commit()

            logger.info(f"Starting ingestion for document: {doc.name} ({doc.format})")

            path = Path(file_path)
            if not path.exists():
                raise FileNotFoundError(f"File not found: {file_path}")

            raw_sections = await extract_content(path, doc.format)
            if not raw_sections:
                raise ValueError("No content could be extracted from document")

            doc.page_count = raw_sections[-1].get("page_number", 1) if raw_sections else 0
            doc.word_count = sum(len(s.get("content", "").split()) for s in raw_sections)
            doc.language = detect_language(raw_sections[0].get("content", "") if raw_sections else "")

            db_sections = []
            for idx, sec_data in enumerate(raw_sections):
                section = DocumentSection(
                    id=str(uuid.uuid4()),
                    document_id=doc_id,
                    title=sec_data.get("title"),
                    content=sec_data.get("content", ""),
                    section_type=sec_data.get("section_type", "body"),
                    level=sec_data.get("level", 1),
                    page_number=sec_data.get("page_number"),
                    sequence_index=idx,
                    word_count=len(sec_data.get("content", "").split()),
                )
                db.add(section)
                db_sections.append(section)

            await db.flush()
            logger.info(f"Saved {len(db_sections)} sections for {doc.name}")

            from app.core.llm_client import get_llm_client
            llm = get_llm_client()

            all_units = []
            for section in db_sections[:20]:
                if len(section.content) < 50:
                    continue
                try:
                    units_data = await llm.extract_knowledge_units(section.content, section.title or "Untitled")
                    for ud in units_data:
                        unit = KnowledgeUnit(
                            id=str(uuid.uuid4()),
                            document_id=doc_id,
                            section_id=section.id,
                            unit_type=ud.get("type", "fact"),
                            content=ud.get("content", ""),
                            title=ud.get("title"),
                            confidence_score=ud.get("confidence", 0.7),
                            tags=ud.get("tags", []),
                        )
                        db.add(unit)
                        all_units.append(unit)
                except Exception as e:
                    logger.warning(f"Knowledge extraction failed for section {section.id}: {e}")

            await db.flush()
            logger.info(f"Extracted {len(all_units)} knowledge units for {doc.name}")

            full_text = " ".join(s.content for s in db_sections[:5])
            try:
                entities_data = await llm.extract_entities(full_text[:3000])
                for ed in entities_data:
                    entity = ExtractedEntity(
                        id=str(uuid.uuid4()),
                        document_id=doc_id,
                        entity_type=ed.get("type", "unknown"),
                        value=ed.get("value", ""),
                        confidence=ed.get("confidence", 0.7),
                        context=ed.get("context"),
                    )
                    db.add(entity)
            except Exception as e:
                logger.warning(f"Entity extraction failed: {e}")

            await db.flush()

            await _build_embeddings(doc_id, db_sections, all_units)

            doc.status = DocumentStatus.PROCESSED
            doc.processing_completed_at = datetime.utcnow()
            await db.commit()

            logger.info(f"Ingestion complete for: {doc.name}")

            from app.api.v1.websocket import get_connection_manager
            ws_manager = get_connection_manager()
            await ws_manager.broadcast(f"doc:{doc_id}", {
                "type": "processing_complete",
                "document_id": doc_id,
                "sections": len(db_sections),
                "knowledge_units": len(all_units),
            })

        except Exception as e:
            logger.exception(f"Ingestion failed for document {doc_id}: {e}")
            async with AsyncSessionLocal() as err_db:
                err_result = await err_db.execute(select(Document).where(Document.id == doc_id))
                err_doc = err_result.scalar_one_or_none()
                if err_doc:
                    err_doc.status = DocumentStatus.FAILED
                    err_doc.processing_error = str(e)
                    await err_db.commit()


async def extract_content(path: Path, format: str) -> List[Dict[str, Any]]:
    extractors = {
        "pdf": extract_pdf,
        "docx": extract_docx,
        "pptx": extract_pptx,
        "csv": extract_csv,
        "xlsx": extract_xlsx,
        "txt": extract_txt,
        "png": extract_image_ocr,
        "jpg": extract_image_ocr,
        "jpeg": extract_image_ocr,
        "tiff": extract_image_ocr,
    }
    extractor = extractors.get(format, extract_txt)
    return await asyncio.get_event_loop().run_in_executor(None, extractor, path)


def extract_pdf(path: Path) -> List[Dict[str, Any]]:
    sections = []
    try:
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                tables = page.extract_tables()
                if text.strip():
                    sections.append({
                        "title": f"Page {page_num}",
                        "content": text.strip(),
                        "section_type": "body",
                        "level": 1,
                        "page_number": page_num,
                    })
                for table in tables:
                    if table:
                        table_text = "\n".join(" | ".join(str(cell or "") for cell in row) for row in table if row)
                        sections.append({
                            "title": f"Table on Page {page_num}",
                            "content": table_text,
                            "section_type": "table",
                            "level": 2,
                            "page_number": page_num,
                        })
    except Exception as e:
        logger.warning(f"pdfplumber failed, trying PyPDF2: {e}")
        try:
            import PyPDF2
            with open(path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page_num, page in enumerate(reader.pages, 1):
                    text = page.extract_text() or ""
                    if text.strip():
                        sections.append({
                            "title": f"Page {page_num}",
                            "content": text.strip(),
                            "section_type": "body",
                            "level": 1,
                            "page_number": page_num,
                        })
        except Exception as e2:
            logger.error(f"PDF extraction completely failed: {e2}")
            if settings.ocr_enabled:
                sections = extract_image_ocr(path)

    return sections or [{"title": "Document", "content": "Could not extract content from PDF.", "section_type": "body", "level": 1, "page_number": 1}]


def extract_docx(path: Path) -> List[Dict[str, Any]]:
    sections = []
    try:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        current_section = None
        current_content = []

        for para in doc.paragraphs:
            if not para.text.strip():
                continue
            if para.style.name.startswith("Heading"):
                if current_section and current_content:
                    sections.append({
                        "title": current_section,
                        "content": "\n".join(current_content),
                        "section_type": "section",
                        "level": int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1,
                        "page_number": None,
                    })
                current_section = para.text.strip()
                current_content = []
            else:
                current_content.append(para.text.strip())

        if current_section and current_content:
            sections.append({
                "title": current_section,
                "content": "\n".join(current_content),
                "section_type": "section",
                "level": 1,
                "page_number": None,
            })
        elif current_content:
            sections.append({
                "title": "Document Body",
                "content": "\n".join(current_content),
                "section_type": "body",
                "level": 1,
                "page_number": None,
            })

        for idx, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            if rows:
                sections.append({
                    "title": f"Table {idx + 1}",
                    "content": "\n".join(rows),
                    "section_type": "table",
                    "level": 2,
                    "page_number": None,
                })
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")

    return sections or [{"title": "Document", "content": "Content extraction failed.", "section_type": "body", "level": 1, "page_number": None}]


def extract_pptx(path: Path) -> List[Dict[str, Any]]:
    sections = []
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_title = None
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if not text:
                        continue
                    if shape.shape_type == 13 or (slide_title is None and shape.name.lower().startswith("title")):
                        slide_title = text
                    else:
                        slide_content.append(text)
            if slide_title or slide_content:
                sections.append({
                    "title": slide_title or f"Slide {slide_num}",
                    "content": "\n".join(slide_content) or slide_title or "",
                    "section_type": "slide",
                    "level": 1,
                    "page_number": slide_num,
                })
    except Exception as e:
        logger.error(f"PPTX extraction failed: {e}")
    return sections


def extract_csv(path: Path) -> List[Dict[str, Any]]:
    try:
        import pandas as pd
        df = pd.read_csv(str(path))
        preview = df.head(50).to_string()
        stats = df.describe(include="all").to_string()
        columns = ", ".join(df.columns.tolist())
        return [
            {
                "title": "Data Overview",
                "content": f"Columns: {columns}\n\nShape: {df.shape[0]} rows × {df.shape[1]} columns\n\nPreview:\n{preview}",
                "section_type": "data",
                "level": 1,
                "page_number": None,
            },
            {
                "title": "Statistical Summary",
                "content": stats,
                "section_type": "table",
                "level": 2,
                "page_number": None,
            },
        ]
    except Exception as e:
        logger.error(f"CSV extraction failed: {e}")
        return []


def extract_xlsx(path: Path) -> List[Dict[str, Any]]:
    sections = []
    try:
        import pandas as pd
        xl = pd.ExcelFile(str(path))
        for sheet_name in xl.sheet_names:
            df = xl.parse(sheet_name)
            preview = df.head(30).to_string()
            sections.append({
                "title": f"Sheet: {sheet_name}",
                "content": f"Shape: {df.shape[0]} rows × {df.shape[1]} columns\n\n{preview}",
                "section_type": "data",
                "level": 1,
                "page_number": None,
            })
    except Exception as e:
        logger.error(f"XLSX extraction failed: {e}")
    return sections


def extract_txt(path: Path) -> List[Dict[str, Any]]:
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        chunks = []
        lines = content.split("\n")
        chunk_size = 100
        for i in range(0, len(lines), chunk_size):
            chunk_lines = lines[i:i + chunk_size]
            chunk_text = "\n".join(chunk_lines).strip()
            if chunk_text:
                chunks.append({
                    "title": f"Section {i // chunk_size + 1}",
                    "content": chunk_text,
                    "section_type": "body",
                    "level": 1,
                    "page_number": None,
                })
        return chunks
    except Exception as e:
        logger.error(f"TXT extraction failed: {e}")
        return []


def extract_image_ocr(path: Path) -> List[Dict[str, Any]]:
    if not settings.ocr_enabled:
        return [{"title": "Image", "content": "OCR is disabled.", "section_type": "body", "level": 1, "page_number": 1}]
    try:
        import pytesseract
        from PIL import Image
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        img = Image.open(str(path))
        text = pytesseract.image_to_string(img)
        if text.strip():
            return [{"title": "OCR Extracted Text", "content": text.strip(), "section_type": "body", "level": 1, "page_number": 1}]
    except Exception as e:
        logger.error(f"OCR extraction failed: {e}")
    return [{"title": "Image", "content": "Could not extract text from image.", "section_type": "body", "level": 1, "page_number": 1}]


def detect_language(text: str) -> str:
    if not text:
        return "en"
    try:
        from langdetect import detect
        return detect(text[:500])
    except Exception:
        return "en"


async def _build_embeddings(doc_id: str, sections: List[DocumentSection], units: List[KnowledgeUnit]):
    try:
        from app.core.embeddings import EmbeddingService
        from app.core.vector_store import VectorStoreService
        import numpy as np

        embedding_svc = EmbeddingService.get()
        vector_store = VectorStoreService.get()

        texts = []
        metadata_list = []

        for unit in units:
            texts.append(unit.content)
            metadata_list.append({
                "unit_id": unit.id,
                "document_id": doc_id,
                "unit_type": unit.unit_type,
                "section_id": unit.section_id,
                "page_number": None,
                "section_title": None,
            })

        if texts:
            embeddings = await asyncio.get_event_loop().run_in_executor(
                None, lambda: embedding_svc.embed(texts)
            )
            vector_store.add_vectors(embeddings, metadata_list)
            logger.info(f"Added {len(texts)} embeddings for document {doc_id}")
    except Exception as e:
        logger.error(f"Embedding pipeline failed for {doc_id}: {e}")
